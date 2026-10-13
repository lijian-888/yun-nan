import asyncio
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from unittest.mock import AsyncMock

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.worksheet.cell_range import CellRange

from app.ai_gateway import AIProviderSettings
from app.research_agent import ResearchAgentError
from app.teacher_file_query import TeacherFileQuery, query_teacher_records, evidence_for_result, validate_teacher_answer, build_teacher_file_evidence
from app.teacher_source_files import create_teacher_source_files_router
from app.teacher_workbook_index import build_index, metric_code, _unit
from app.teacher_workbook_reader import _number_text


class TeacherFileQueryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name) / "sources"
        self.folder = self.root / "黄平老师"
        self.folder.mkdir(parents=True)
        self.file = self.folder / "2023勐海考种.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "数据汇总"
        sheet.append(["2023编号", "2022编号", "株高(cm)", "直链淀粉2022", "直链淀粉2023", "千粒重", "地点", "重复"])
        sheet.append(["DT1", "C1-83", 96, 0.1408, 0.1222, 34.96, "勐海", "Ⅰ"])
        sheet.append(["DT10", "C1-99", 108, 0.14, 0.12, 32, "勐海", "Ⅰ"])
        sheet.append(["DT1", "C1-83", 98, 0.15, 0.13, 35, "勐海", "Ⅱ"])
        sheet.append(["空值", "", None, "#N/A", "=1+1", None, "勐海", "Ⅰ"])
        sheet["D5"].data_type = "e"
        other = book.create_sheet("多层表头")
        other.append(["材料编号", "株高(cm)", None, "产量(kg/亩)"])
        other.merge_cells("B1:C1")
        other.append([None, "Ⅰ", "Ⅱ", "平均"])
        other.append(["TEST1", 101, 102, 600])
        other["XFD100"].value = None  # Format-only far cells must not expand the index.
        unlabeled = book.create_sheet("无标题材料列")
        unlabeled.append([None, "穗长", "千粒重"])
        unlabeled.append(["UNC1", 23, 28])
        unlabeled.append(["UNC2", 24, 29])
        book.save(self.file)
        self.index = Path(self.temp.name) / "index.sqlite"
        result = build_index(self.root, self.index)
        self.assertFalse(result["errors"])

    def tearDown(self):
        self.temp.cleanup()

    def query(self, **kw):
        return query_teacher_records(TeacherFileQuery(**kw), root=self.root, target=self.index)

    def test_exact_identity_and_receipt(self):
        result = self.query(material="DT10", metric="株高", year=2023)
        self.assertEqual(result["status"], "found")
        row = result["records"][0]
        self.assertEqual((row["value"], row["cell"], row["sheet"], row["unit"]), ("108", "C3", "数据汇总", "cm"))
        context, cards, states = evidence_for_result(result, "DT10的株高")
        self.assertEqual(cards[0]["source_location"]["cell"], "C3")
        validate_teacher_answer("2023年DT10的株高为108厘米。【原表1】", states[-1])
        with self.assertRaises(ResearchAgentError):
            validate_teacher_answer("株高为109厘米。【原表1】", states[-1])

    def test_repeated_records_clarify_without_returning_values(self):
        result = self.query(material="DT1", metric="株高", year=2023)
        self.assertEqual((result["status"], result["count"], result["records"]), ("ambiguous", 2, []))
        self.assertTrue(result["options"])
        self.assertNotIn("value", result["candidates"][0])
        context, cards, states = evidence_for_result(result, "原问题")
        self.assertEqual(states[-1]["kind"], "teacher_file")
        self.assertEqual(cards, [])
        with self.assertRaises(ResearchAgentError):
            validate_teacher_answer("按第一条株高96cm", states[1])

    def test_repeat_filter_cell_selection_and_all_do_not_average(self):
        first = self.query(material="DT1", metric="株高", year=2023, conditions={"重复": "Ⅰ"})
        self.assertEqual(first["records"][0]["value"], "96")
        chosen = self.query(material="DT1", metric="株高", cell="C4", sheet="数据汇总")
        self.assertEqual(chosen["records"][0]["value"], "98")
        all_rows = self.query(material="DT1", metric="株高", show_all=True)
        self.assertEqual([r["value"] for r in all_rows["records"]], ["96", "98"])
        _, _, states = evidence_for_result(all_rows, "查看全部")
        with self.assertRaises(ResearchAgentError):
            validate_teacher_answer("平均株高97cm。【原表1】", states[-1])
        with self.assertRaises(ResearchAgentError):
            validate_teacher_answer("株高98cm。【原表1】\n株高96cm。【原表2】", states[-1])

    def test_column_year_overrides_file_year_and_does_not_invent_percent(self):
        result = self.query(material="DT10", metric="直链淀粉含量", year=2022)
        self.assertEqual(result["records"][0]["cell"], "D3")
        self.assertEqual(result["records"][0]["value"], "0.14")
        self.assertEqual(result["records"][0]["year_basis"], "指标表头")
        self.assertEqual(result["records"][0]["unit"], "")
        _, _, states = evidence_for_result(result, "原问题")
        with self.assertRaises(ResearchAgentError):
            validate_teacher_answer("直链淀粉14%。【原表1】", states[-1])

    def test_missing_formula_and_error_not_zero(self):
        empty = self.query(material="空值", metric="株高")
        self.assertEqual(empty["records"][0]["status"], "missing")
        error = self.query(material="空值", metric="直链淀粉", year=2022)
        self.assertEqual(error["records"][0]["status"], "error")
        formula = self.query(material="空值", metric="直链淀粉", year=2023)
        self.assertEqual(formula["records"][0]["status"], "formula_uncached")
        self.assertEqual(formula["records"][0]["value"], "")

    def test_merged_multilevel_headers_keep_two_repetitions(self):
        result = self.query(material="TEST1", metric="株高")
        self.assertEqual(result["status"], "ambiguous")
        self.assertEqual([r["cell"] for r in result["candidates"]], ["B3", "C3"])

    def test_exact_not_fuzzy_and_filename_years_not_all_available(self):
        result = self.query(material="DT100", metric="株高")
        self.assertEqual(result["status"], "not_found")
        self.assertEqual(result["records"], [])
        wrong_year = self.query(material="DT10", metric="株高", year=2022)
        self.assertEqual(wrong_year["status"], "not_found")
        self.assertEqual(wrong_year["available_years"], [2023])

    def test_changed_or_new_sources_fail_closed(self):
        self.file.write_bytes(self.file.read_bytes() + b"changed")
        with self.assertRaisesRegex(ValueError, "目录已变化"):
            self.query(material="DT10", metric="株高")

    def test_authenticated_api_and_metadata_path_boundary(self):
        with patch.dict("os.environ", {"YNAAS_TEACHER_FILES_DIR": str(self.root), "YNAAS_TEACHER_INDEX_PATH": str(self.index)}):
            app = FastAPI()
            app.include_router(create_teacher_source_files_router(lambda: None))
            client = TestClient(app)
            result = client.post("/api/research/teacher-source-files/query", json={"material":"DT10","metric":"株高","year":2023})
            self.assertEqual(result.status_code, 200)
            self.assertEqual(result.json()["records"][0]["cell"], "C3")
            self.assertEqual(client.get("/api/research/teacher-source-files/metadata", params={"path":"../escape.xlsx"}).status_code, 400)
            self.assertEqual(client.get("/api/research/teacher-source-files/metadata", params={"path":"黄平老师/2023勐海考种.xlsx", "expected_sha256":"0"*64}).status_code, 409)
            def denied():
                raise HTTPException(403, "无权限")
            blocked = FastAPI()
            blocked.include_router(create_teacher_source_files_router(denied))
            self.assertEqual(TestClient(blocked).get("/api/research/teacher-source-files/catalog").status_code, 403)

    def test_narrative_not_header_and_trait_distinctions(self):
        self.assertIsNone(metric_code("株高158cm，生育期168天，易落粒"))
        self.assertNotEqual(metric_code("实粒数"), metric_code("每穗实粒数"))
        result = self.query(material="DT10", metric="千粒重")
        self.assertEqual(result["records"][0]["unit"], "")

    def test_model_parameters_and_privacy_policy(self):
        provider = AIProviderSettings("cherryin", "https://example.invalid", "model", "test", True)
        with patch.dict("os.environ", {"ALLOW_SHARED_BUSINESS_EGRESS": "false"}):
            with self.assertRaisesRegex(ResearchAgentError, "未获准"):
                asyncio.run(build_teacher_file_evidence(provider=provider, question="查询", history=[]))

    def test_numeric_material_display_and_compound_value_are_preserved(self):
        self.assertEqual(_number_text("358.0", "General"), "358")
        self.assertEqual(_number_text("7", "000"), "007")
        scope = {"status":"found", "records":[{"value":"75/115", "display_value":"75/115 cm", "status":"source_value", "citation":"【原表1】"}]}
        validate_teacher_answer("株高原记录75/115厘米。【原表1】", scope)
        with self.assertRaises(ResearchAgentError):
            validate_teacher_answer("株高115厘米。【原表1】", scope)

    def test_unlabeled_material_column_requires_explicit_confirmation(self):
        result = self.query(material="UNC1", metric="穗长")
        self.assertEqual(result["status"], "ambiguous")
        self.assertTrue(result["identity_pending"])
        self.assertEqual(result["records"], [])
        _, _, states = evidence_for_result(result, "原问题")
        self.assertFalse(states[-1]["allow_skip"])
        confirmed = self.query(material="UNC1", metric="穗长", confirm_identity=True, sheet="无标题材料列", cell="B2")
        self.assertEqual(confirmed["records"][0]["value"], "23")

    def test_planner_retries_once_without_guessed_parameters(self):
        provider = AIProviderSettings("cherryin", "https://example.invalid", "model", "test", True)
        planner = AsyncMock(side_effect=[ResearchAgentError("JSON failed"), {"material":"DT10","metric":"株高","year":2023}])
        with patch.dict("os.environ", {"ALLOW_SHARED_BUSINESS_EGRESS":"true","YNAAS_TEACHER_FILES_DIR":str(self.root),"YNAAS_TEACHER_INDEX_PATH":str(self.index)}), patch("app.teacher_file_query.model_json_request", planner):
            _, _, states = asyncio.run(build_teacher_file_evidence(provider=provider,question="DT10在2023年的株高",history=[]))
        self.assertEqual(planner.await_count, 2)
        self.assertEqual(states[1]["records"][0]["value"], "108")

    def test_single_source_citation_and_per_area_unit(self):
        self.assertEqual(_unit("试验面积（亩）", "0.04"), "亩")
        self.assertEqual(_unit("浓度（mg/kg）", "1.2"), "mg/kg")
        self.assertEqual(_unit("面积（m²）", "20"), "m2")
        scope={"status":"found","records":[{"value":"600","display_value":"600 kg/亩","unit":"kg/亩","status":"source_value","citation":"【原表1】"}]}
        validate_teacher_answer("产量600公斤/亩。\n来源：【原表1】",scope)
        with self.assertRaises(ResearchAgentError):
            validate_teacher_answer("产量600公斤。【原表1】",scope)
